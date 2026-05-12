from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

COLUMNS_TO_COMPARE = [
    'POLIZA', 'PLACA', 'MARCA', 'LÍNEA', 'TOMADOR', 'NOMBRE TOMADOR',
    'ASEGURADO', 'NOMBRE ASEGURADO', 'FECHA INICIO VIGE', 'FECHA VENC',
    'VALOR ASEGURADO 2025', 'PRIMA NETA 2025', 'PRIMA TOTAL 2025'
]

COLUMNS_ALL = [
    'POLIZA', 'PLACA', 'MARCA', 'LÍNEA', 'TOMADOR', 'NOMBRE TOMADOR',
    'ASEGURADO', 'NOMBRE ASEGURADO', 'CELULAR', 'CORREO', 'NIT ONEROSO',
    'BENEFICIARIO ONEROSO', 'FECHA INICIO VIGE', 'FECHA VENC',
    'VALOR ASEGURADO 2024', 'VALOR ASEGURADO 2025', 'SINIESTROS',
    'PRIMA NETA 2024', 'PRIMA TOTAL 2024', 'PRIMA NETA 2025', 'PRIMA TOTAL 2025',
    'VARIACIÓN PRIMA NETA', 'VARIACIÓN PRIMA TOTAL'
]

class ExcelService:
    def __init__(self):
        self.columns = COLUMNS_ALL

    def _normalize_header(self, header):
        if header is None:
            return ''
        return str(header).strip().upper().replace(' ', '').replace('.', '').replace('_', '').replace('Ó', 'O').replace('Í', 'I')

    def _normalize_value(self, val):
        if val is None:
            return ''
        if isinstance(val, float):
            if val == int(val):
                return str(int(val))
            return str(val)
        if isinstance(val, (int,)):
            return str(val)
        return str(val).strip()

    def create_workbook(self, data_list):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Pólizas Bolívar'

        self._apply_header_style(ws, self.columns)

        for row_idx, data in enumerate(data_list, 2):
            for col_idx, col_name in enumerate(self.columns, 1):
                value = data.get(col_name, '')
                if 'FECHA' in col_name.upper() and isinstance(value, datetime):
                    value = value.strftime('%d/%m/%Y')
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = self._get_border()
                cell.alignment = Alignment(horizontal='left', vertical='center')

        self._adjust_column_widths(ws, self.columns)

        return wb

    def create_comparison_workbook(self, extracted_data, insurer_data):
        wb = Workbook()

        ws_extract = wb.active
        ws_extract.title = 'Datos Extraídos'

        ws_comparison = wb.create_sheet('Comparativo')

        self._build_extracted_sheet(ws_extract, extracted_data)

        self._build_vertical_comparison_sheet(ws_comparison, extracted_data, insurer_data)

        return wb

    def _build_extracted_sheet(self, ws, extracted_data):
        self._apply_header_style(ws, COLUMNS_ALL)

        for row_idx, data in enumerate(extracted_data, 2):
            for col_idx, col_name in enumerate(COLUMNS_ALL, 1):
                value = data.get(col_name, '')
                if 'FECHA' in col_name.upper():
                    value = self._format_value(value, col_name)
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = self._get_border()
                cell.alignment = Alignment(horizontal='left', vertical='center')

        self._adjust_column_widths(ws, COLUMNS_ALL)

    def _build_vertical_comparison_sheet(self, ws, extracted_data, insurer_data):
        header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True, size=11)
        thin_border = self._get_border()

        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 35
        ws.column_dimensions['C'].width = 35
        ws.column_dimensions['D'].width = 12

        ws.cell(row=1, column=1, value='CAMPO').fill = header_fill
        ws.cell(row=1, column=1).font = header_font
        ws.cell(row=1, column=2, value='VALOR PDF').fill = header_fill
        ws.cell(row=1, column=2).font = header_font
        ws.cell(row=1, column=3, value='VALOR EXCEL ASEGURADORA').fill = header_fill
        ws.cell(row=1, column=3).font = header_font
        ws.cell(row=1, column=4, value='OK').fill = header_fill
        ws.cell(row=1, column=4).font = header_font

        ws.cell(row=1, column=1).border = thin_border
        ws.cell(row=1, column=2).border = thin_border
        ws.cell(row=1, column=3).border = thin_border
        ws.cell(row=1, column=4).border = thin_border

        ws.row_dimensions[1].height = 25

        header_map = self._build_insurer_header_map(insurer_data)

        current_row = 2

        for ext_data in extracted_data:
            poliza = str(ext_data.get('POLIZA', '')).strip()
            placa = str(ext_data.get('PLACA', '')).strip()

            ref_data = self._find_insurer_row(insurer_data, header_map, poliza, placa)

            archivo = ext_data.get('ARCHIVO', '')

            section_fill = PatternFill(start_color='D6DCE4', end_color='D6DCE4', fill_type='solid')
            section_font = Font(bold=True, size=11, color='1F4E79')

            ws.cell(row=current_row, column=1, value=f'📄 {archivo}').fill = section_fill
            ws.cell(row=current_row, column=1).font = section_font
            ws.cell(row=current_row, column=1).alignment = Alignment(horizontal='left', vertical='center')
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
            ws.row_dimensions[current_row].height = 22
            current_row += 1

            if ref_data:
                placa_ref = str(ref_data.get('PLACA', '')).strip()
                ws.cell(row=current_row, column=1, value=f'   Póliza: {poliza} | Placa: {placa_ref}').fill = section_fill
            else:
                ws.cell(row=current_row, column=1, value=f'   Póliza: {poliza} | Placa: {placa} (sin coincidencia en Excel)').fill = section_fill
            ws.cell(row=current_row, column=1).font = Font(italic=True, size=10)
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
            ws.row_dimensions[current_row].height = 18
            current_row += 1

            for col_name in COLUMNS_TO_COMPARE:
                ext_val = ext_data.get(col_name, '')
                ref_val = ''
                if ref_data:
                    ref_val = self._get_insurer_value(ref_data, header_map, col_name)
                    if ref_val is None:
                        ref_val = ''

                match = self._compare_values(ext_val, ref_val)

                label_cell = ws.cell(row=current_row, column=1, value=col_name)
                label_cell.border = thin_border
                label_cell.alignment = Alignment(horizontal='left', vertical='center')

                ext_display = self._format_value(ext_val, col_name)
                pdf_cell = ws.cell(row=current_row, column=2, value=ext_display)
                pdf_cell.border = thin_border
                pdf_cell.alignment = Alignment(horizontal='left', vertical='center')

                ref_display = self._format_value(ref_val, col_name)
                excel_cell = ws.cell(row=current_row, column=3, value=ref_display)
                excel_cell.border = thin_border
                excel_cell.alignment = Alignment(horizontal='left', vertical='center')

                ok_cell = ws.cell(row=current_row, column=4, value=match)
                ok_cell.border = thin_border
                ok_cell.alignment = Alignment(horizontal='center', vertical='center')

                if match == '✓':
                    ok_cell.fill = PatternFill(start_color='CCFFCC', end_color='CCFFCC', fill_type='solid')
                elif match == '✗':
                    ok_cell.fill = PatternFill(start_color='FFCCCC', end_color='FFCCCC', fill_type='solid')
                    pdf_cell.fill = PatternFill(start_color='FFF3CD', end_color='FFF3CD', fill_type='solid')
                    excel_cell.fill = PatternFill(start_color='FFF3CD', end_color='FFF3CD', fill_type='solid')
                elif match == '-':
                    ok_cell.fill = PatternFill(start_color='E0E0E0', end_color='E0E0E0', fill_type='solid')

                current_row += 1

            current_row += 1

    def _format_value(self, val, col_name):
        if val is None or val == '':
            return ''

        if 'FECHA' in col_name.upper():
            val_str = str(val)
            if isinstance(val, datetime):
                return val.strftime('%d/%m/%Y')

            for fmt in ['%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d.%m.%Y']:
                try:
                    return datetime.strptime(val_str.split()[0] if ' ' in val_str else val_str, fmt).strftime('%d/%m/%Y')
                except:
                    pass
            return val_str

        if any(x in col_name.upper() for x in ['VALOR', 'PRIMA']):
            if isinstance(val, (int, float)):
                return f'{val:,.0f}'.replace(',', '.')

        return str(val)

    def _build_insurer_header_map(self, insurer_data):
        if not insurer_data:
            return {}
        result = {}
        for key in insurer_data[0].keys():
            norm_key = self._normalize_header(str(key))
            result[norm_key] = key
        return result

    def _find_insurer_row(self, insurer_data, header_map, poliza, placa):
        poliza_norm = self._normalize_value(poliza).replace('.', '').replace(',', '').replace(' ', '')

        poliza_variants = ['POLIZA', 'NOPOLIZA', 'POLIZANO']
        placa_variants = ['PLACA', 'PLACAVEHI']

        for row in insurer_data:
            for variant in poliza_variants:
                norm_variant = self._normalize_header(variant)
                if norm_variant in header_map:
                    val = row.get(header_map[norm_variant], '')
                    if val is not None:
                        val_str = self._normalize_value(val).replace('.', '').replace(',', '').replace(' ', '')
                        if poliza_norm and val_str and (poliza_norm in val_str or val_str in poliza_norm):
                            return row

            for variant in placa_variants:
                norm_variant = self._normalize_header(variant)
                if norm_variant in header_map:
                    val = row.get(header_map[norm_variant], '')
                    if val and placa and str(placa).upper().strip() == str(val).upper().strip():
                        return row
        return None

    def _get_insurer_value(self, insurer_row, header_map, target_col):
        header_mappings = {
            'POLIZA': ['POLIZA', 'NOPOLIZA', 'POLIZANO'],
            'PLACA': ['PLACA', 'PLACAVEHI'],
            'MARCA': ['MARCA'],
            'LÍNEA': ['LÍNEA', 'LINEA', 'CLASE'],
            'TOMADOR': ['TOMADOR', 'IDTOMADOR'],
            'NOMBRE TOMADOR': ['NOMBRETOMADOR', 'NOMBRE TOMADOR'],
            'ASEGURADO': ['ASEGURADO', 'IDASEGURADO'],
            'NOMBRE ASEGURADO': ['NOMBREASEGURADO', 'NOMBRE ASEGURADO'],
            'FECHA INICIO VIGE': ['FECHAINICIOVIGE', 'FECHA INICIO VIGE'],
            'FECHA VENC': ['FECHAVENC', 'FECHA VENC'],
            'VALOR ASEGURADO 2025': ['VALORASEGURADO2025', 'VALOR ASEGURADO 2025'],
            'PRIMA NETA 2025': ['PRIMANETA2025', 'PRIMA NETA 2025'],
            'PRIMA TOTAL 2025': ['PRIMATOTAL2025', 'PRIMA TOTAL 2025'],
        }

        variants = header_mappings.get(target_col, [target_col])
        for variant in variants:
            norm_variant = self._normalize_header(variant)
            if norm_variant in header_map:
                return insurer_row.get(header_map[norm_variant])
        return None

    def _compare_values(self, val1, val2):
        s1 = self._normalize_value(val1).upper().strip()
        s2 = self._normalize_value(val2).upper().strip()

        if s2 in ['#N/D', 'N/D', '', 'NONE', '-', 'N/A']:
            return '-'

        if s1 == s2:
            return '✓'

        date_formats = ['%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d.%m.%Y']
        val1_date = None
        val2_date = None

        if '/' in s1 or '-' in s1 or '.' in s1:
            for fmt in date_formats:
                try:
                    val1_date = datetime.strptime(s1, fmt)
                    break
                except:
                    continue

        for fmt in date_formats:
            try:
                val2_date = datetime.strptime(s2, fmt)
                break
            except:
                continue

        if val1_date and val2_date:
            if val1_date == val2_date:
                return '✓'

        try:
            n1 = float(s1.replace('$', '').replace(',', '').replace('.', ''))
            n2 = float(s2.replace('$', '').replace(',', '').replace('.', ''))
            if abs(n1 - n2) < 100:
                return '✓'
        except:
            pass

        return '✗'

    def load_insurer_excel(self, filepath):
        try:
            wb = load_workbook(filepath, data_only=True)
            ws = wb.active

            headers = [cell.value for cell in ws[1]]

            data = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is None:
                    continue
                row_dict = {}
                for i, header in enumerate(headers):
                    if header is not None:
                        row_dict[header] = row[i] if i < len(row) else None
                data.append(row_dict)

            return data
        except Exception as e:
            raise Exception(f'Error al leer Excel de aseguradora: {str(e)}')

    def calculate_kpis(self, extracted_data, insurer_data):
        kpis = {
            'total_procesados': len(extracted_data),
            'coincidencias_exactas': 0,
            'diferencias_encontradas': 0,
            'sin_coincidencia_excel': 0,
            'con_errores': 0,
            'resumen_campos': []
        }

        header_map = self._build_insurer_header_map(insurer_data)

        detalle_campos = {col: {'ok': 0, 'diferencia': 0, 'sin_datos': 0} for col in COLUMNS_TO_COMPARE}

        for data in extracted_data:
            if 'ERROR' in data:
                kpis['con_errores'] += 1
                continue

            poliza = str(data.get('POLIZA', '')).strip()
            placa = str(data.get('PLACA', '')).strip()

            ref_data = self._find_insurer_row(insurer_data, header_map, poliza, placa)

            if not ref_data:
                kpis['sin_coincidencia_excel'] += 1
                continue

            tiene_diferencia = False
            for col in COLUMNS_TO_COMPARE:
                ext_val = data.get(col, '')
                ref_val = self._get_insurer_value(ref_data, header_map, col) if ref_data else ''
                if ref_val is None:
                    ref_val = ''

                match = self._compare_values(ext_val, ref_val)

                if match == '✓':
                    detalle_campos[col]['ok'] += 1
                elif match == '✗':
                    detalle_campos[col]['diferencia'] += 1
                    tiene_diferencia = True
                else:
                    detalle_campos[col]['sin_datos'] += 1

            if tiene_diferencia:
                kpis['diferencias_encontradas'] += 1
            else:
                kpis['coincidencias_exactas'] += 1

        for col in COLUMNS_TO_COMPARE:
            total = sum(detalle_campos[col].values())
            ok = detalle_campos[col]['ok']
            dif = detalle_campos[col]['diferencia']
            sin = detalle_campos[col]['sin_datos']

            kpis['resumen_campos'].append({
                'campo': col,
                'ok': ok,
                'diferencia': dif,
                'sin_datos': sin,
                'total': total,
                'porcentaje_ok': round((ok / total * 100), 1) if total > 0 else 0,
                'porcentaje_dif': round((dif / total * 100), 1) if total > 0 else 0
            })

        return kpis

    def _apply_header_style(self, ws, columns):
        header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True, size=10)
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        for col_idx, col_name in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = self._get_border()

        ws.row_dimensions[1].height = 35

    def _get_border(self):
        return Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

    def _adjust_column_widths(self, ws, columns):
        for col_idx in range(1, len(columns) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 18

    def save_workbook(self, wb, filepath):
        wb.save(filepath)
